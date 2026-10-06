"""Build stride-analysis.xlsx and risk-register.xlsx from data.py."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import data

ROOT = Path(__file__).resolve().parent.parent
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="CBD2D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
RATING_FILL = {"Low": "C6EFCE", "Medium": "FFEB9C", "High": "F8CBAD", "Critical": "FF8B8B"}


def sheet(wb, title, headers, rows, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.fill, c.font, c.alignment, c.border = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical="center"), BORDER
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment, c.border = WRAP, BORDER
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def rating_formula(cell):
    return f'=IF({cell}>=17,"Critical",IF({cell}>=10,"High",IF({cell}>=5,"Medium","Low")))'


def colour_ratings(ws, col_letter, n):
    rng = f"{col_letter}2:{col_letter}{n + 1}"
    for name, fill in RATING_FILL.items():
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="equal", formula=[f'"{name}"'], fill=PatternFill("solid", fgColor=fill)))


def notes_sheet(wb, title, lines):
    ws = wb.active
    ws.title = title
    for line in lines:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 120
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True)


def stride_workbook():
    wb = Workbook()
    a = data.ASSESSMENT
    notes_sheet(wb, "README", [
        f"{a['title']} - STRIDE analysis",
        f"Client: {a['client']} (fictional)   Version {a['version']}   {a['date']}   Author: {a['author']}",
        "",
        "Sheets:",
        "  Zones / Trust Boundaries / Assets / Data Flows - the system model the threats are derived from.",
        "  STRIDE per Element - which STRIDE categories were considered for each element type (Microsoft STRIDE-per-element).",
        "  Threat Register - 30 threats with actor, path, existing controls, scoring, mitigation and residual risk.",
        "  Actors - threat actor catalogue referenced in the register.",
        "",
        "Scoring: Likelihood (1-5) x Impact (1-5). Score/rating columns are live formulas.",
        "Full methodology: risk-assessment/risk-methodology.md. Source of truth: tools/data.py.",
    ])
    sheet(wb, "Zones", ["Zone", "Name", "Purdue", "SL-T (indicative)", "SL-A (indicative)", "Description"],
          [[z["id"], z["name"], z["purdue"], z["sl_t"], z["sl_a"], z["desc"]] for z in data.ZONES],
          [8, 26, 14, 16, 16, 70])
    sheet(wb, "Trust Boundaries", ["ID", "Between", "Enforced by", "State", "Observation"],
          [[t["id"], t["between"], t["enforced_by"], t["state"], t["observation"]] for t in data.TRUST_BOUNDARIES],
          [8, 36, 32, 18, 80])
    sheet(wb, "Assets", ["ID", "Asset", "Zone", "Purdue", "Criticality", "Function", "Security concern", "Threats"],
          [[x["id"], x["name"], x["zone"], x["purdue"], x["criticality"], x["function"], x["concern"],
            ", ".join(t["id"] for t in data.THREATS if x["id"] in t["assets"])] for x in data.ASSETS],
          [8, 34, 7, 8, 12, 44, 40, 30])
    sheet(wb, "Data Flows", ["ID", "Source", "Destination", "Protocol", "Trust boundary", "Description", "Threats"],
          [[f["id"], f["src"], f["dst"], f["proto"], f["tb"], f["desc"],
            ", ".join(t["id"] for t in data.THREATS if f["id"] in t["flows"])] for f in data.DATA_FLOWS],
          [8, 24, 24, 32, 20, 44, 24])
    matrix = [
        ["External entity", "X", "", "X", "", "", ""],
        ["Process", "X", "X", "X", "X", "X", "X"],
        ["Data store", "", "X", "(if logs)", "X", "X", ""],
        ["Data flow", "", "X", "", "X", "X", ""],
    ]
    sheet(wb, "STRIDE per Element", ["Element type", "S", "T", "R", "I", "D", "E"], matrix, [20, 8, 8, 10, 8, 8, 8])

    headers = ["Threat ID", "Title", "STRIDE", "Category", "Assets", "Data flows", "Threat description",
               "Threat actor", "Attack path", "Existing controls", "Likelihood", "Impact", "Score", "Rating",
               "Recommended mitigation", "Residual L", "Residual I", "Residual score", "Residual rating",
               "Attack scenarios"]
    rows = []
    for i, t in enumerate(data.enriched_threats(), start=2):
        rows.append([t["id"], t["title"], t["stride"], t["stride_name"], ", ".join(t["assets"]),
                     ", ".join(t["flows"]), t["desc"], t["actor"], t["path"], t["existing"], t["L"], t["I"],
                     f"=K{i}*L{i}", rating_formula(f"M{i}"),
                     "; ".join(f"{c} {data.control(c)['name']}" for c in t["controls"]),
                     t["rL"], t["rI"], f"=P{i}*Q{i}", rating_formula(f"R{i}"), ", ".join(t["scenarios"])])
    ws = sheet(wb, "Threat Register", headers, rows,
               [9, 32, 7, 18, 16, 14, 60, 12, 40, 36, 10, 8, 7, 10, 48, 9, 9, 9, 10, 14])
    colour_ratings(ws, "N", len(rows))
    colour_ratings(ws, "S", len(rows))
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=False)
    ws.add_data_validation(dv)
    for col in "KLPQ":
        dv.add(f"{col}2:{col}{len(rows) + 1}")
    sheet(wb, "Actors", ["ID", "Threat actor"], [[k, v] for k, v in data.ACTORS.items()], [8, 50])
    out = ROOT / "threat-model" / "stride-analysis.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print("wrote", out.relative_to(ROOT))


def risk_workbook():
    wb = Workbook()
    notes_sheet(wb, "Methodology", [
        "Risk methodology - Apex Manufacturing (fictional)",
        "Risk score = Likelihood x Impact (5 x 5). This is the assessment's documented methodology, not a universal standard.",
        "Current risk = with existing controls. Residual risk = after the recommended controls are implemented.",
        "Thresholds: 1-4 Low | 5-9 Medium | 10-16 High | 17-25 Critical",
        "Top-10 ranking: score desc, then impact desc, then number of attack scenarios the threat appears in, then ID.",
        "",
        "LIKELIHOOD",
        *[f"  {n} {name}: {d}" for n, name, d in data.LIKELIHOOD_SCALE],
        "",
        "IMPACT (highest applicable dimension: safety, production, financial, quality, data)",
        *[f"  {n} {name}: {d}" for n, name, d in data.IMPACT_SCALE],
    ])
    ts = data.enriched_threats()
    rows = []
    for i, t in enumerate(ts, start=2):
        rows.append([t["id"], t["title"], t["stride_name"], ", ".join(data.asset_name(a) for a in t["assets"]),
                     t["L"], t["I"], f"=E{i}*F{i}", rating_formula(f"G{i}"), ", ".join(t["controls"]),
                     t["rL"], t["rI"], f"=J{i}*K{i}", rating_formula(f"L{i}"), f"=G{i}-L{i}"])
    ws = sheet(wb, "Risk Register",
               ["Threat ID", "Risk", "STRIDE", "Assets", "L", "I", "Score", "Rating", "Controls",
                "Res. L", "Res. I", "Res. score", "Res. rating", "Reduction"],
               rows, [9, 40, 20, 40, 5, 5, 7, 10, 24, 7, 7, 9, 11, 10])
    colour_ratings(ws, "H", len(rows))
    colour_ratings(ws, "M", len(rows))
    n = len(rows) + 1

    # Heat maps driven by COUNTIFS over the register
    hm = wb.create_sheet("Heat Map")
    for block, (lc, ic, title, top) in enumerate([("E", "F", "CURRENT RISK", 1), ("J", "K", "RESIDUAL RISK", 10)]):
        hm.cell(top, 1, f"{title} - number of threats (rows = likelihood, columns = impact)").font = Font(bold=True)
        for imp in range(1, 6):
            hm.cell(top + 1, imp + 1, f"I{imp}").font = Font(bold=True)
        for li in range(5, 0, -1):
            r = top + 2 + (5 - li)
            hm.cell(r, 1, f"L{li}").font = Font(bold=True)
            for imp in range(1, 6):
                c = hm.cell(r, imp + 1, f"=COUNTIFS('Risk Register'!{lc}2:{lc}{n},{li},'Risk Register'!{ic}2:{ic}{n},{imp})")
                c.fill = PatternFill("solid", fgColor=RATING_FILL[data.rating(li * imp)])
                c.alignment = Alignment(horizontal="center")
                c.border = BORDER
    hm.column_dimensions["A"].width = 8

    top = data.top_risks()
    sheet(wb, "Top 10", ["Rank", "Threat ID", "Risk", "Score", "Rating", "Scenarios", "Residual", "Key controls"],
          [[k, t["id"], t["title"], t["score"], t["rating"], ", ".join(t["scenarios"]),
            f"{t['rscore']} {t['rrating']}", ", ".join(t["controls"])] for k, t in enumerate(top, 1)],
          [6, 9, 46, 7, 10, 18, 12, 22])
    sheet(wb, "Controls Mapping",
          ["Control", "Name", "Description", "IEC 62443", "NIST SP 800-53 (800-82r3 overlay)", "NIST CSF 2.0", "Threats mitigated"],
          [[c["id"], c["name"], c["desc"], c["iec"], c["nist"], c["csf"], ", ".join(data.threats_for_control(c["id"]))]
           for c in data.CONTROLS], [8, 30, 70, 24, 26, 20, 30])
    sheet(wb, "Roadmap", ["ID", "Horizon", "Action", "Controls", "Threats addressed", "Owner", "Effort"],
          [[r["id"], r["horizon"], r["action"], ", ".join(r["controls"]),
            ", ".join(sorted({t for c in r["controls"] for t in data.threats_for_control(c)})),
            r["owner"], r["effort"]] for r in data.ROADMAP], [7, 12, 80, 16, 36, 22, 9])
    out = ROOT / "risk-assessment" / "risk-register.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print("wrote", out.relative_to(ROOT))


if __name__ == "__main__":
    data.validate()
    stride_workbook()
    risk_workbook()
